from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json
import os
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from backend.ipm_update_store import (
    IPMUpdateError,
    LOCAL_TZ,
    _as_datetime,
    _clean,
    _compare_records,
    _find_header,
    _load_bootstrap,
    _load_index as _legacy_load_index,
    _read_gzip_json as _legacy_read_gzip_json,
    _parse_protocol,
    _value,
    parse_ipm_xlsx,
    validate_ipm_xlsx,
)


EDGE_BASE = os.getenv(
    "SUPABASE_IPM_EDGE_URL",
    "https://mccjtfgdbhtkuqwjlnbg.supabase.co/functions/v1/seplanbi-ipm",
)

OPTIONAL_ALIASES = {
    "applicant_name": ("Requerente", "Requerente - Nome Razão", "Requerente - Nome/Razão"),
    "applicant_document": ("RequerenteCPF/CNPJ", "Requerente - CPF/CNPJ"),
    "opening_note": ("ObsAbertura", "Abertura - Observação"),
    "last_movement_note": ("UltTramiteObs", "Último Trâmite - Observação"),
    "last_activity": ("ÚltimaAtividade", "Última Atividade"),
    "responsible_name": ("Responsável", "Responsável - Nome"),
    "responsible_document": ("ResponsávelCPF/CNPJ", "Responsável - CPF/CNPJ"),
    "opening_sector": ("CCAbertura", "Centro de Custo Abertura - Descrição"),
    "source_current_user": ("UsuárioAtual", "Usuário Atual"),
}


def _edge(
    action: str,
    *,
    payload: dict[str, Any] | None = None,
    body: bytes | None = None,
    headers: dict[str, str] | None = None,
    query: dict[str, str] | None = None,
    expect_binary: bool = False,
    timeout: int = 25,
):
    try:
        from vercel.oidc import get_vercel_oidc_token_sync
        token = str(get_vercel_oidc_token_sync() or "").strip()
    except Exception:
        token = os.getenv("VERCEL_OIDC_TOKEN", "").strip()
    if not token:
        raise IPMUpdateError(503, "A conexão segura Vercel → Supabase não está disponível.")

    params = {"action": action}
    if query:
        params.update({key: str(value) for key, value in query.items()})
    url = f"{EDGE_BASE}?{urlencode(params)}"

    request_headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/octet-stream" if expect_binary else "application/json",
    }
    request_headers.update(headers or {})

    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        request_headers["Content-Type"] = "application/json"
    else:
        data = body

    request = Request(url, data=data, headers=request_headers, method="POST" if data is not None else "GET")
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read()
    except HTTPError as exc:
        raw = exc.read()
        try:
            message = json.loads(raw.decode("utf-8")).get("error")
        except Exception:
            message = None
        raise IPMUpdateError(
            exc.code if 400 <= exc.code <= 599 else 503,
            message or "O Supabase recusou a atualização da base IPM.",
        ) from exc
    except (URLError, TimeoutError) as exc:
        raise IPMUpdateError(503, "Não foi possível acessar o armazenamento Supabase.") from exc

    if expect_binary:
        return raw
    try:
        return json.loads(raw.decode("utf-8")) if raw else {}
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IPMUpdateError(503, "O Supabase retornou uma resposta inválida.") from exc


def _iso_local(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=LOCAL_TZ)
    return value.isoformat(timespec="seconds")


def _optional_positions(headers: list[str]) -> dict[str, int]:
    positions: dict[str, int] = {}
    for field, aliases in OPTIONAL_ALIASES.items():
        for alias in aliases:
            if alias in headers:
                positions[field] = headers.index(alias)
                break
    return positions


def extract_stage_rows(body: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise IPMUpdateError(400, "O arquivo não é uma planilha XLSX válida.") from exc

    try:
        if "Report" not in workbook.sheetnames:
            raise IPMUpdateError(400, "A aba Report não foi encontrada no relatório IPM.")
        sheet = workbook["Report"]
        positions, header_row, headers = _find_header(sheet)
        positions.update(_optional_positions(headers))

        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(_clean(value) for value in values):
                continue
            protocol = _parse_protocol(_value(values, positions.get("protocol")))
            if not protocol:
                continue

            opened = _as_datetime(_value(values, positions.get("opened")))
            moved = _as_datetime(_value(values, positions.get("moved")))
            closed = _as_datetime(_value(values, positions.get("closed")))

            rows.append({
                "protocol_id": protocol,
                "opened_at": _iso_local(opened),
                "last_movement_at": _iso_local(moved),
                "closed_at": _iso_local(closed),
                "source_status": _clean(_value(values, positions.get("situation"))),
                "subject": _clean(_value(values, positions.get("subject"))),
                "current_sector": _clean(_value(values, positions.get("sector"))),
                "applicant_name": _clean(_value(values, positions.get("applicant_name"))) or None,
                "applicant_document": _clean(_value(values, positions.get("applicant_document"))) or None,
                "opening_note": _clean(_value(values, positions.get("opening_note"))) or None,
                "last_movement_note": _clean(_value(values, positions.get("last_movement_note"))) or None,
                "last_activity": _clean(_value(values, positions.get("last_activity"))) or None,
                "responsible_name": _clean(_value(values, positions.get("responsible_name"))) or None,
                "responsible_document": _clean(_value(values, positions.get("responsible_document"))) or None,
                "opening_sector": _clean(_value(values, positions.get("opening_sector"))) or None,
                "source_current_user": _clean(_value(values, positions.get("source_current_user"))) or None,
                "diff_status": "unchanged",
                "diff_fields": {},
                "validation_errors": [],
            })
        return rows
    finally:
        workbook.close()


def _canonical_name(reference_date: str) -> str:
    date = datetime.strptime(reference_date, "%Y-%m-%d")
    return f"BASEIPM_{date.strftime('%d%m%Y')}.xlsx"


def _steps(status: str, staged_rows: int = 0, source_rows: int = 0) -> list[dict[str, str]]:
    ready = status in {"ready_for_review", "approved", "rejected"}
    approved = status == "approved"
    failed = status == "failed"
    staged = source_rows > 0 and staged_rows == source_rows
    return [
        {"id": "received", "label": "Arquivo recebido", "status": "done"},
        {"id": "stored", "label": "Original armazenado no Supabase", "status": "done" if status != "uploaded" else "pending"},
        {"id": "validated", "label": "Estrutura e integridade validadas", "status": "done" if not failed else "blocked"},
        {"id": "snapshot", "label": "Base carregada em staging", "status": "done" if staged else ("blocked" if failed else "pending")},
        {"id": "comparison", "label": "Comparação com a base anterior", "status": "done" if ready or failed else "pending"},
        {"id": "events", "label": "Validação automática da atualização", "status": "done" if ready else ("blocked" if failed else "pending")},
        {"id": "review", "label": "Validação humana da base", "status": "done" if approved else ("blocked" if status == "rejected" else "pending")},
        {"id": "indicators", "label": "Recálculo dos indicadores", "status": "blocked"},
        {"id": "publication", "label": "Publicação do BI", "status": "blocked"},
    ]


def _ui_run(raw: dict[str, Any], issues: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    status = str(raw.get("status") or "validating")
    source_rows = int(raw.get("source_rows") or 0)
    staged_rows = int(raw.get("staged_rows") or 0)
    reference_date = str(raw.get("reference_date") or "")
    metadata = raw.get("source_metadata") or {}
    validation = raw.get("validation_summary") or {}

    if status == "ready_for_review":
        ui_status, phase = "REVIEW", "AWAITING_APPROVAL"
        message = "Validação automática concluída. Revise os registros e aprove ou rejeite a base."
    elif status == "approved":
        ui_status, phase = "DONE", "APPROVED"
        message = "Base aprovada. A publicação do BI permanece bloqueada até a etapa de indicadores."
    elif status == "rejected":
        ui_status, phase = "REJECTED", "REJECTED"
        message = raw.get("review_notes") or "Base rejeitada na validação humana."
    elif status == "failed":
        ui_status, phase = "FAILED", "VALIDATION_FAILED"
        message = raw.get("review_notes") or "A base falhou em um ou mais gates de validação."
    else:
        ui_status, phase = "PROCESSING", "VALIDATION_COMPLETE"
        message = "Arquivo armazenado. Preparando staging e comparação no Supabase."

    comparison = {
        "baseline_rows": max(0, source_rows - int(raw.get("new_count") or 0) + int(raw.get("removed_count") or 0)),
        "current_rows": source_rows,
        "new": int(raw.get("new_count") or 0),
        "changed": int(raw.get("changed_count") or 0),
        "unchanged": int(raw.get("unchanged_count") or 0),
        "removed": int(raw.get("removed_count") or 0),
        "field_changes": metadata.get("field_changes") or {
            "opened": 0,
            "last_movement": 0,
            "closed": 0,
            "situation": 0,
            "subject": 0,
            "sector": 0,
        },
        "event_counts": metadata.get("event_counts") or {},
    }

    return {
        "id": str(raw.get("import_run_id") or raw.get("id") or ""),
        "status": ui_status,
        "db_status": status,
        "phase": phase,
        "created_at": raw.get("created_at"),
        "local_created_at": raw.get("created_at"),
        "source_name": raw.get("source_filename") or "",
        "canonical_name": metadata.get("canonical_name") or (_canonical_name(reference_date) if reference_date else ""),
        "metrics": {
            "rows": source_rows,
            "unique_protocols": source_rows,
            "duplicates": int(raw.get("duplicate_count") or 0),
            "latest_movement": metadata.get("latest_movement") or reference_date,
            "source_sheet": metadata.get("source_sheet") or "Report",
            "source_columns": int(metadata.get("source_columns") or 0),
            "date_sentinels_ignored": int(metadata.get("date_sentinels_ignored") or 0),
            "warnings": int(raw.get("warning_count") or validation.get("warnings") or 0),
            "errors": int(raw.get("error_count") or validation.get("errors") or 0),
            "comparison": comparison,
        },
        "steps": _steps(status, staged_rows, source_rows),
        "can_approve": bool(raw.get("can_approve")),
        "can_publish": False,
        "issues": issues or [],
        "message": message,
    }


def create_ipm_import(body: bytes, source_name: str = "") -> dict[str, Any]:
    if not _clean(source_name).lower().endswith(".xlsx"):
        raise IPMUpdateError(400, "Selecione um arquivo XLSX do relatório IPM.")

    metrics = validate_ipm_xlsx(body)
    reference_date = str(metrics["latest_movement"])[:10]
    canonical_name = _canonical_name(reference_date)

    created = _edge(
        "create-run",
        payload={
            "reference_date": reference_date,
            "source_filename": _clean(source_name)[:240],
            "source_rows": int(metrics["unique_protocols"]),
            "duplicate_count": int(metrics["duplicates"]),
            "invalid_date_count": 0,
            "source_metadata": {
                **metrics,
                "canonical_name": canonical_name,
            },
        },
    )
    run = created.get("run")
    if not isinstance(run, dict) or not run.get("id"):
        raise IPMUpdateError(503, "O Supabase não criou o registro da atualização.")

    run_id = str(run["id"])
    try:
        _edge(
            "upload-file",
            body=body,
            headers={
                "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "X-Run-Id": run_id,
                "X-Reference-Date": reference_date,
                "X-Canonical-Name": canonical_name,
            },
        )
    except Exception:
        try:
            _edge("fail", payload={"run_id": run_id, "reason": "Falha ao armazenar o XLSX original no Supabase."})
        except Exception:
            pass
        raise

    run.update({
        "import_run_id": run_id,
        "staged_rows": 0,
        "source_metadata": {**metrics, "canonical_name": canonical_name},
    })
    return {"ok": True, "duplicate": False, "run": _ui_run(run)}


LEGACY_BOOTSTRAP_ROWS = 3170


def _load_initial_baseline() -> tuple[dict[str, dict[str, str]], dict[str, Any]]:
    """Carrega o snapshot auditado anterior para a primeira migração ao Supabase.

    Prioriza o snapshot persistido no Vercel Blob legado. O arquivo local do
    repositório é apenas fallback para ambientes de desenvolvimento e precisa
    ser um GZIP válido.
    """
    legacy_error: Exception | None = None
    try:
        for item in _legacy_load_index():
            snapshot_path = str(item.get("snapshot_path") or "").strip()
            if not snapshot_path:
                continue
            metrics = item.get("metrics") or {}
            declared_rows = int(metrics.get("unique_protocols") or metrics.get("rows") or 0)
            if declared_rows and declared_rows != LEGACY_BOOTSTRAP_ROWS:
                continue
            payload = _legacy_read_gzip_json(snapshot_path)
            records = payload.get("records") if isinstance(payload, dict) else None
            if isinstance(records, dict) and len(records) == LEGACY_BOOTSTRAP_ROWS:
                return records, {
                    "kind": "legacy_blob_snapshot",
                    "source_rows": len(records),
                    "source_latest_movement": metrics.get("latest_movement"),
                    "snapshot_path": snapshot_path,
                }
    except Exception as exc:
        legacy_error = exc

    try:
        records, metadata = _load_bootstrap()
        if len(records) != LEGACY_BOOTSTRAP_ROWS:
            raise IPMUpdateError(
                503,
                f"Baseline local possui {len(records)} protocolos; esperado: {LEGACY_BOOTSTRAP_ROWS}.",
            )
        return records, metadata
    except Exception as exc:
        if legacy_error is not None:
            raise IPMUpdateError(
                503,
                "Baseline inicial auditado de 3.170 protocolos não foi localizado no armazenamento legado.",
            ) from legacy_error
        if isinstance(exc, IPMUpdateError):
            raise
        raise IPMUpdateError(503, "Baseline inicial da atualização IPM está indisponível.") from exc


def _bootstrap_classification(
    body: bytes,
    stage_rows: list[dict[str, Any]],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    metrics, current_records = parse_ipm_xlsx(body, include_records=True)
    previous_records, _ = _load_initial_baseline()
    comparison, events = _compare_records(previous_records, current_records)

    event_map = {str(item.get("protocol")): item for item in events if item.get("protocol")}
    removed: list[dict[str, Any]] = []

    for row in stage_rows:
        event = event_map.get(row["protocol_id"])
        if not event:
            row["diff_status"] = "unchanged"
            continue
        types = set(event.get("types") or [])
        if "NOVO_PROTOCOLO" in types:
            row["diff_status"] = "new"
        else:
            row["diff_status"] = "changed"
            row["diff_fields"] = event.get("fields") or {}

    for item in events:
        if "AUSENTE_NA_FONTE" not in set(item.get("types") or []):
            continue
        protocol = str(item.get("protocol"))
        previous = previous_records.get(protocol, {})
        removed.append({
            "protocol_id": protocol,
            "details": {
                "opened_at": previous.get("o"),
                "last_movement_at": previous.get("m"),
                "closed_at": previous.get("c"),
                "source_status": previous.get("s"),
                "subject": previous.get("u"),
                "current_sector": previous.get("r"),
            },
        })

    if (
        metrics["unique_protocols"] == 3310
        and metrics["latest_movement"] == "2026-09-28 15:29:43"
    ):
        expected = {"new": 140, "changed": 246, "unchanged": 2924, "removed": 0}
        actual = {key: int(comparison[key]) for key in expected}
        if actual != expected:
            raise IPMUpdateError(
                409,
                "A comparação com a base de referência divergiu do lote auditado. A atualização foi bloqueada.",
            )

    return comparison, removed


def process_ipm_import(run_id: str) -> dict[str, Any]:
    status_payload = _edge("status", query={"id": run_id})
    raw_run = status_payload.get("run") or {}
    if raw_run.get("status") in {"ready_for_review", "approved", "rejected"}:
        return {"ok": True, "run": _ui_run(raw_run, status_payload.get("issues") or []), "already_processed": True}

    body = _edge("download-file", payload={"run_id": run_id}, expect_binary=True)
    stage_rows = extract_stage_rows(body)

    # Primeiro persiste o staging bruto. Assim uma falha no bootstrap/comparação
    # não deixa a execução sem evidência do que foi efetivamente recebido.
    if int(raw_run.get("staged_rows") or 0) != len(stage_rows):
        for offset in range(0, len(stage_rows), 400):
            _edge("stage", payload={"run_id": run_id, "rows": stage_rows[offset:offset + 400]})

    baseline = _edge("baseline-info")
    protocol_count = int(baseline.get("protocol_count") or 0)

    bootstrap_comparison: dict[str, Any] | None = None
    removed_rows: list[dict[str, Any]] = []
    if protocol_count == 0:
        bootstrap_comparison, removed_rows = _bootstrap_classification(body, stage_rows)

        # O bootstrap atribui new/changed/unchanged em memória; reaplica os lotes
        # por upsert para que o staging persistido reflita a comparação auditada.
        for offset in range(0, len(stage_rows), 400):
            _edge("stage", payload={"run_id": run_id, "rows": stage_rows[offset:offset + 400]})

    for offset in range(0, len(removed_rows), 400):
        _edge("stage-removed", payload={"run_id": run_id, "rows": removed_rows[offset:offset + 400]})

    if protocol_count > 0:
        classified = _edge("classify", payload={"run_id": run_id}).get("result") or {}
    else:
        classified = bootstrap_comparison or {}

    validated = _edge("validate", payload={"run_id": run_id}).get("result") or {}
    status_payload = _edge("status", query={"id": run_id})
    raw_run = status_payload.get("run") or {}

    metadata = dict(raw_run.get("source_metadata") or {})
    metadata["field_changes"] = (bootstrap_comparison or {}).get("field_changes", {})
    metadata["event_counts"] = (bootstrap_comparison or {}).get("event_counts", {})
    raw_run["source_metadata"] = metadata

    # The database owns the authoritative counts; this only enriches the UI response.
    raw_run.update({
        "new_count": validated.get("new", classified.get("new", raw_run.get("new_count", 0))),
        "changed_count": validated.get("changed", classified.get("changed", raw_run.get("changed_count", 0))),
        "unchanged_count": validated.get("unchanged", classified.get("unchanged", raw_run.get("unchanged_count", 0))),
        "removed_count": validated.get("removed", classified.get("removed", raw_run.get("removed_count", 0))),
    })

    return {
        "ok": True,
        "run": _ui_run(raw_run, status_payload.get("issues") or []),
        "already_processed": False,
    }


def process_next_ipm_import() -> dict[str, Any]:
    """Uma unidade de trabalho; o banco detém o lease e limita concorrência."""
    claim = _edge("claim-next").get("result") or {}
    run_id = claim.get("run_id")
    lease_token = claim.get("lease_token")
    if not run_id or not lease_token:
        return {"ok": True, "processed": False}

    try:
        result = process_ipm_import(str(run_id))
    except Exception as exc:
        # O erro fica registrado para a próxima tentativa, sem promover dados.
        try:
            _edge("finish-worker", payload={
                "run_id": run_id, "lease_token": lease_token,
                "error": exc.public_message if isinstance(exc, IPMUpdateError) else "Falha interna no processamento.",
            })
        except Exception:
            pass  # O lease expira e permite retomada após falha de rede.
        raise

    _edge("finish-worker", payload={"run_id": run_id, "lease_token": lease_token})
    return {"ok": True, "processed": True, "run_id": str(run_id), "status": result["run"]["db_status"]}


def list_ipm_imports(limit: int = 10) -> dict[str, Any]:
    payload = _edge("history", query={"limit": str(max(1, min(30, int(limit))))})
    runs = [_ui_run(item) for item in payload.get("runs") or []]
    return {"ok": True, "runs": runs, "count": len(runs)}


def get_ipm_import(run_id: str) -> dict[str, Any]:
    payload = _edge("status", query={"id": run_id})
    return {"ok": True, "run": _ui_run(payload.get("run") or {}, payload.get("issues") or [])}


def list_ipm_review_rows(
    run_id: str,
    *,
    diff: str = "",
    q: str = "",
    limit: int = 100,
    offset: int = 0,
) -> dict[str, Any]:
    payload = _edge(
        "review-rows",
        payload={
            "run_id": run_id,
            "diff": diff,
            "q": q,
            "limit": max(1, min(200, int(limit))),
            "offset": max(0, int(offset)),
        },
    )
    return payload


def approve_ipm_import(run_id: str) -> dict[str, Any]:
    _edge("approve", payload={"run_id": run_id})
    return get_ipm_import(run_id)


def reject_ipm_import(run_id: str, reason: str = "") -> dict[str, Any]:
    _edge("reject", payload={"run_id": run_id, "reason": reason})
    return get_ipm_import(run_id)


def supabase_ipm_health() -> dict[str, Any]:
    return _edge("ping")


__all__ = [
    "IPMUpdateError",
    "approve_ipm_import",
    "create_ipm_import",
    "get_ipm_import",
    "list_ipm_imports",
    "list_ipm_review_rows",
    "process_ipm_import",
    "reject_ipm_import",
    "supabase_ipm_health",
]
