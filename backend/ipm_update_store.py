from __future__ import annotations

from collections import Counter
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
import gzip
import json
import os
from pathlib import Path
from typing import Any
import unicodedata
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW_PREFIX = "ipm/raw/"
RUN_PREFIX = "ipm/runs/"
RUN_INDEX_PATH = "ipm/runs/index.json"
SNAPSHOT_PREFIX = "ipm/snapshots/"
EVENT_PREFIX = "ipm/events/"
BOOTSTRAP_FILE = ROOT / "data" / "ipm_bootstrap_v01.json.gz"
MAX_HISTORY = 30
LOCAL_TZ = ZoneInfo("America/Sao_Paulo")

HEADER_ALIASES = {
    "protocol": ("Número/Ano", "Protocolo"),
    "opened": ("Abertura - Data", "DataAbertura"),
    "moved": ("Último Trâmite - Data/Hora", "DataUltTramite", "UltTramiteData"),
    "closed": ("Data Encerramento", "DataEncerramento"),
    "situation": ("Situação",),
    "subject": ("Subassunto - Descrição", "Assunto", "Categoria"),
    "sector": ("Centro de Custo Atual - Descrição", "CCAtual"),
}
REQUIRED_FIELDS = {"protocol", "opened", "moved", "situation", "subject", "sector"}
SNAPSHOT_FIELDS = ("o", "m", "c", "s", "u", "r")


class IPMUpdateError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.public_message = message


def _blob_client():
    try:
        from vercel.blob import BlobClient
    except ImportError as exc:
        raise IPMUpdateError(503, "Armazenamento da atualização IPM indisponível.") from exc
    if not os.getenv("BLOB_READ_WRITE_TOKEN"):
        raise IPMUpdateError(503, "Armazenamento da atualização IPM não está configurado.")
    return BlobClient()


def _read_blob(path: str) -> bytes | None:
    try:
        from vercel.blob.errors import BlobNotFoundError
        with _blob_client() as client:
            result = client.get(path, access="private", timeout=8, use_cache=False)
    except BlobNotFoundError:
        return None
    except IPMUpdateError:
        raise
    except Exception as exc:
        raise IPMUpdateError(503, "Não foi possível consultar o armazenamento da atualização.") from exc
    if result is None or result.status_code == 404:
        return None
    if result.status_code != 200:
        raise IPMUpdateError(503, "O armazenamento recusou a leitura da atualização.")
    return bytes(result.content)


def _write_blob(path: str, body: bytes, content_type: str, *, overwrite: bool) -> None:
    try:
        with _blob_client() as client:
            client.put(path, body, access="private", content_type=content_type, overwrite=overwrite)
    except IPMUpdateError:
        raise
    except Exception as exc:
        raise IPMUpdateError(503, "Não foi possível salvar a atualização no armazenamento privado.") from exc


def _json_blob(path: str, default):
    raw = _read_blob(path)
    if raw is None:
        return default
    try:
        return json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IPMUpdateError(503, "Histórico de atualização IPM inválido.") from exc


def _write_json(path: str, payload: dict, *, overwrite: bool = True) -> None:
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    _write_blob(path, body, "application/json; charset=utf-8", overwrite=overwrite)


def _read_gzip_json(path: str) -> dict | None:
    raw = _read_blob(path)
    if raw is None:
        return None
    try:
        return json.loads(gzip.decompress(raw).decode("utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IPMUpdateError(503, "Snapshot diário armazenado está inválido.") from exc


def _write_gzip_json(path: str, payload: dict, *, overwrite: bool = True) -> None:
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    _write_blob(path, gzip.compress(raw, compresslevel=9), "application/gzip", overwrite=overwrite)


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _norm_text(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", _clean(value)).split()).casefold()


def _parse_protocol(value: Any) -> str:
    text = _clean(value)
    sep = "/" if "/" in text else "-" if "-" in text else ""
    if not sep:
        return ""
    left, right = text.rsplit(sep, 1)
    try:
        a, b = int(left.strip()), int(right.strip())
    except ValueError:
        return ""
    if sep == "/":
        number, year = a, b
    elif 2000 <= a <= 2100:
        year, number = a, b
    else:
        number, year = a, b
    if not (2025 <= year <= 2100) or number <= 0:
        return ""
    return f"{year}-{number}"


def _as_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)):
        numeric = float(value)
        if numeric < 30000:
            return None
        return datetime(1899, 12, 30) + timedelta(days=numeric)
    text = _clean(value)
    if not text:
        return None
    try:
        numeric = float(text.replace(",", "."))
        if numeric < 30000:
            return None
        return datetime(1899, 12, 30) + timedelta(days=numeric)
    except ValueError:
        pass
    for fmt in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            continue
    return None


def _dt_text(value: datetime | None) -> str:
    return "" if value is None else value.strftime("%Y-%m-%dT%H:%M:%S")


def _numeric_sentinel(value: Any) -> bool:
    text = _clean(value)
    if not text:
        return False
    try:
        return float(text.replace(",", ".")) < 30000
    except ValueError:
        return False


def _find_header(sheet) -> tuple[dict[str, int], int, list[str]]:
    for row_no, values in enumerate(sheet.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
        headers = [_clean(value) for value in values]
        positions: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                if alias in headers:
                    positions[field] = headers.index(alias)
                    break
        if REQUIRED_FIELDS.issubset(positions):
            return positions, row_no, headers
    raise IPMUpdateError(
        400,
        "Estrutura do relatório não reconhecida. Os campos essenciais do relatório IPM não foram encontrados.",
    )


def _value(values: tuple, index: int | None):
    if index is None or index >= len(values):
        return None
    return values[index]


def parse_ipm_xlsx(body: bytes, *, include_records: bool = False) -> tuple[dict[str, Any], dict[str, dict[str, str]]]:
    if not body:
        raise IPMUpdateError(400, "O arquivo enviado está vazio.")
    try:
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise IPMUpdateError(400, "O arquivo não é uma planilha XLSX válida.") from exc

    try:
        if "Report" not in workbook.sheetnames:
            raise IPMUpdateError(400, "A aba Report não foi encontrada no relatório IPM.")
        sheet = workbook["Report"]
        positions, header_row, headers = _find_header(sheet)

        seen: set[str] = set()
        duplicates: list[str] = []
        invalid_rows = 0
        invalid_dates = 0
        sentinels_ignored = 0
        latest_movement: datetime | None = None
        rows = 0
        records: dict[str, dict[str, str]] = {}

        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(_clean(value) for value in values):
                continue
            protocol = _parse_protocol(_value(values, positions.get("protocol")))
            if not protocol:
                invalid_rows += 1
                continue
            rows += 1
            if protocol in seen:
                duplicates.append(protocol)
                continue
            seen.add(protocol)

            opened_raw = _value(values, positions.get("opened"))
            moved_raw = _value(values, positions.get("moved"))
            closed_raw = _value(values, positions.get("closed")) if "closed" in positions else None

            opened = _as_datetime(opened_raw)
            moved = _as_datetime(moved_raw)
            closed = _as_datetime(closed_raw)

            if opened is None or moved is None or moved < opened:
                invalid_dates += 1
                continue
            if closed is not None and closed < opened:
                invalid_dates += 1
                continue
            if closed is None and _numeric_sentinel(closed_raw):
                sentinels_ignored += 1
            if latest_movement is None or moved > latest_movement:
                latest_movement = moved

            if include_records:
                records[protocol] = {
                    "o": _dt_text(opened),
                    "m": _dt_text(moved),
                    "c": _dt_text(closed),
                    "s": _norm_text(_value(values, positions.get("situation"))),
                    "u": _norm_text(_value(values, positions.get("subject"))),
                    "r": _norm_text(_value(values, positions.get("sector"))),
                }

        if invalid_rows:
            raise IPMUpdateError(400, f"O relatório contém {invalid_rows} linha(s) com protocolo inválido.")
        if duplicates:
            raise IPMUpdateError(400, f"O relatório contém {len(duplicates)} protocolo(s) duplicado(s).")
        if invalid_dates:
            raise IPMUpdateError(400, f"O relatório contém {invalid_dates} inconsistência(s) de data.")
        if not seen or latest_movement is None:
            raise IPMUpdateError(400, "Nenhum protocolo válido foi encontrado no relatório.")

        metrics = {
            "rows": rows,
            "unique_protocols": len(seen),
            "duplicates": 0,
            "latest_movement": latest_movement.isoformat(sep=" ", timespec="seconds"),
            "source_sheet": "Report",
            "source_columns": len([header for header in headers if header]),
            "date_sentinels_ignored": sentinels_ignored,
        }
        return metrics, records
    finally:
        workbook.close()


def validate_ipm_xlsx(body: bytes) -> dict[str, Any]:
    metrics, _ = parse_ipm_xlsx(body, include_records=False)
    return metrics


def _steps() -> list[dict[str, str]]:
    return [
        {"id": "received", "label": "Arquivo recebido", "status": "done"},
        {"id": "stored", "label": "Original armazenado", "status": "done"},
        {"id": "validated", "label": "Estrutura e integridade validadas", "status": "done"},
        {"id": "snapshot", "label": "Snapshot diário registrado", "status": "done"},
        {"id": "comparison", "label": "Comparação com a base anterior", "status": "pending"},
        {"id": "events", "label": "Identificação de movimentações", "status": "pending"},
        {"id": "indicators", "label": "Recálculo dos indicadores", "status": "pending"},
        {"id": "preview", "label": "Validação em pré-produção", "status": "pending"},
        {"id": "publication", "label": "Publicação do BI", "status": "blocked"},
    ]


def _step(run: dict, step_id: str, status: str) -> None:
    for item in run.get("steps", []):
        if item.get("id") == step_id:
            item["status"] = status
            return


def _load_index() -> list[dict]:
    payload = _json_blob(RUN_INDEX_PATH, {"v": 2, "runs": []})
    runs = payload.get("runs", []) if isinstance(payload, dict) else []
    return runs if isinstance(runs, list) else []


def _save_index(runs: list[dict]) -> None:
    _write_json(RUN_INDEX_PATH, {"v": 2, "runs": runs[:MAX_HISTORY]}, overwrite=True)


def _summary(run: dict) -> dict:
    return {
        "id": run["id"],
        "status": run["status"],
        "phase": run["phase"],
        "created_at": run["created_at"],
        "local_created_at": run["local_created_at"],
        "source_name": run.get("source_name", ""),
        "canonical_name": run["canonical_name"],
        "metrics": run.get("metrics", {}),
        "steps": run.get("steps", []),
        "snapshot_path": run.get("snapshot_path"),
        "events_path": run.get("events_path"),
        "can_publish": bool(run.get("can_publish", False)),
        "message": run.get("message", ""),
    }


def _persist_run(run: dict) -> None:
    _write_json(f"{RUN_PREFIX}{run['id']}.json", run, overwrite=True)
    history = [item for item in _load_index() if item.get("id") != run["id"]]
    _save_index([_summary(run), *history])


def list_ipm_imports(limit: int = 10) -> dict[str, Any]:
    limit = max(1, min(30, int(limit)))
    runs = _load_index()[:limit]
    return {"ok": True, "runs": runs, "count": len(runs)}


def get_ipm_import(run_id: str) -> dict[str, Any]:
    safe = "".join(ch for ch in _clean(run_id) if ch.isalnum() or ch in "-_")
    if not safe or safe != _clean(run_id):
        raise IPMUpdateError(400, "Identificador de atualização inválido.")
    payload = _json_blob(f"{RUN_PREFIX}{safe}.json", None)
    if not isinstance(payload, dict):
        raise IPMUpdateError(404, "Atualização IPM não encontrada.")
    return {"ok": True, "run": payload}


def _load_bootstrap() -> tuple[dict[str, dict[str, str]], dict[str, Any]]:
    try:
        payload = json.loads(gzip.decompress(BOOTSTRAP_FILE.read_bytes()).decode("utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IPMUpdateError(503, "Baseline inicial da atualização IPM está indisponível.") from exc
    records = payload.get("records")
    if not isinstance(records, dict):
        raise IPMUpdateError(503, "Baseline inicial da atualização IPM está inválida.")
    return records, {
        "kind": "bootstrap",
        "source_rows": int(payload.get("source_rows") or len(records)),
        "source_latest_movement": payload.get("source_latest_movement"),
    }


def _previous_snapshot(run: dict) -> tuple[dict[str, dict[str, str]], dict[str, Any]]:
    target_created = str(run.get("created_at") or "")
    for item in _load_index():
        if item.get("id") == run.get("id"):
            continue
        created = str(item.get("created_at") or "")
        snapshot_path = item.get("snapshot_path")
        if target_created and created and created >= target_created:
            continue
        if not snapshot_path:
            continue
        payload = _read_gzip_json(str(snapshot_path))
        records = payload.get("records") if isinstance(payload, dict) else None
        if isinstance(records, dict):
            return records, {
                "kind": "snapshot",
                "run_id": item.get("id"),
                "canonical_name": item.get("canonical_name"),
                "source_rows": len(records),
                "source_latest_movement": item.get("metrics", {}).get("latest_movement"),
            }
    return _load_bootstrap()


def _compare_records(previous: dict[str, dict[str, str]], current: dict[str, dict[str, str]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    previous_ids = set(previous)
    current_ids = set(current)
    new_ids = sorted(current_ids - previous_ids)
    removed_ids = sorted(previous_ids - current_ids)
    common = previous_ids & current_ids

    field_changes = Counter()
    changed_ids: list[str] = []
    events: list[dict[str, Any]] = []

    for protocol in new_ids:
        events.append({"protocol": protocol, "types": ["NOVO_PROTOCOLO"], "fields": {}})

    for protocol in sorted(common):
        before = previous[protocol]
        after = current[protocol]
        changed_fields = [field for field in SNAPSHOT_FIELDS if before.get(field, "") != after.get(field, "")]
        if not changed_fields:
            continue
        changed_ids.append(protocol)
        types: list[str] = []
        detail: dict[str, dict[str, str]] = {}
        for field in changed_fields:
            field_changes[field] += 1
            detail[field] = {"before": before.get(field, ""), "after": after.get(field, "")}
            if field == "o":
                types.append("ABERTURA_ALTERADA")
            elif field == "m":
                types.append("NOVO_TRAMITE")
            elif field == "s":
                types.append("MUDOU_SITUACAO")
            elif field == "u":
                types.append("MUDOU_ASSUNTO")
            elif field == "r":
                types.append("MUDOU_SETOR")
            elif field == "c":
                if not before.get("c") and after.get("c"):
                    types.append("ENCERRADO")
                elif before.get("c") and not after.get("c"):
                    types.append("REABERTO")
                else:
                    types.append("ENCERRAMENTO_ALTERADO")
        events.append({"protocol": protocol, "types": types, "fields": detail})

    for protocol in removed_ids:
        events.append({"protocol": protocol, "types": ["AUSENTE_NA_FONTE"], "fields": {}})

    event_counts = Counter(event for item in events for event in item["types"])
    comparison = {
        "baseline_rows": len(previous),
        "current_rows": len(current),
        "new": len(new_ids),
        "changed": len(changed_ids),
        "unchanged": len(common) - len(changed_ids),
        "removed": len(removed_ids),
        "field_changes": {
            "opened": field_changes["o"],
            "last_movement": field_changes["m"],
            "closed": field_changes["c"],
            "situation": field_changes["s"],
            "subject": field_changes["u"],
            "sector": field_changes["r"],
        },
        "event_counts": dict(sorted(event_counts.items())),
    }
    return comparison, events


def _raw_path_for_run(run: dict) -> str:
    if run.get("raw_path"):
        return str(run["raw_path"])
    try:
        local_created = datetime.fromisoformat(str(run["local_created_at"]))
    except (KeyError, ValueError) as exc:
        raise IPMUpdateError(503, "Execução sem referência válida para o arquivo original.") from exc
    return f"{RAW_PREFIX}{local_created.strftime('%Y/%m')}/{run['canonical_name']}"


def process_ipm_import(run_id: str) -> dict[str, Any]:
    run = get_ipm_import(run_id)["run"]
    if run.get("phase") == "DIFF_COMPLETE":
        return {"ok": True, "run": run, "already_processed": True}

    run["status"] = "PROCESSING"
    run["phase"] = "COMPARING"
    _step(run, "comparison", "processing")
    run["message"] = "Comparando o snapshot recebido com a última base válida."
    _persist_run(run)

    try:
        raw = _read_blob(_raw_path_for_run(run))
        if raw is None:
            raise IPMUpdateError(503, "Arquivo original desta execução não foi encontrado.")

        current_metrics, current_records = parse_ipm_xlsx(raw, include_records=True)
        if current_metrics["unique_protocols"] != int(run.get("metrics", {}).get("unique_protocols", -1)):
            raise IPMUpdateError(409, "O snapshot reprocessado diverge da validação original.")

        previous_records, baseline = _previous_snapshot(run)
        comparison, events = _compare_records(previous_records, current_records)

        # Gate de referência da primeira migração: reproduz o delta já auditado
        # entre a extração de 17/09 e a extração de 28/09.
        reference_gate = "not_applicable"
        if (
            baseline.get("kind") == "bootstrap"
            and baseline.get("source_rows") == 3170
            and current_metrics["unique_protocols"] == 3310
            and current_metrics["latest_movement"] == "2026-09-28 15:29:43"
        ):
            expected = {"new": 140, "changed": 246, "removed": 0, "unchanged": 2924}
            actual = {key: comparison[key] for key in expected}
            if actual != expected:
                raise IPMUpdateError(409, "A comparação automática divergiu do lote auditado de referência. A pipeline foi bloqueada.")
            reference_gate = "passed"

        snapshot_path = f"{SNAPSHOT_PREFIX}{run['id']}.json.gz"
        events_path = f"{EVENT_PREFIX}{run['id']}.json.gz"
        _write_gzip_json(
            snapshot_path,
            {
                "v": 1,
                "run_id": run["id"],
                "canonical_name": run["canonical_name"],
                "latest_movement": current_metrics["latest_movement"],
                "records": current_records,
            },
            overwrite=True,
        )
        _write_gzip_json(
            events_path,
            {
                "v": 1,
                "run_id": run["id"],
                "baseline": baseline,
                "comparison": comparison,
                "events": events,
            },
            overwrite=True,
        )

        run["snapshot_path"] = snapshot_path
        run["events_path"] = events_path
        run["baseline"] = baseline
        run["metrics"] = {**run.get("metrics", {}), **current_metrics, "comparison": comparison, "reference_gate": reference_gate}
        run["status"] = "PREPARED"
        run["phase"] = "DIFF_COMPLETE"
        _step(run, "comparison", "done")
        _step(run, "events", "done")
        run["can_publish"] = False
        run["message"] = (
            f"Comparação concluída: {comparison['new']} novos, {comparison['changed']} alterados, "
            f"{comparison['unchanged']} sem alteração e {comparison['removed']} ausentes. "
            "O recálculo dos indicadores e a publicação continuam bloqueados nesta versão."
        )
        _persist_run(run)
        return {"ok": True, "run": run, "already_processed": False}
    except IPMUpdateError as exc:
        run["status"] = "REVIEW"
        run["phase"] = "DIFF_FAILED"
        _step(run, "comparison", "blocked")
        run["message"] = exc.public_message
        _persist_run(run)
        raise
    except Exception as exc:
        run["status"] = "FAILED"
        run["phase"] = "DIFF_FAILED"
        _step(run, "comparison", "blocked")
        run["message"] = "Falha técnica durante a comparação. Nenhuma alteração foi publicada."
        _persist_run(run)
        raise IPMUpdateError(500, run["message"]) from exc


def create_ipm_import(body: bytes, source_name: str = "") -> dict[str, Any]:
    if not _clean(source_name).lower().endswith(".xlsx"):
        raise IPMUpdateError(400, "Selecione um arquivo XLSX do relatório IPM.")

    validation = validate_ipm_xlsx(body)
    import hashlib
    digest = hashlib.sha256(body).hexdigest()
    now_utc = datetime.now(timezone.utc).replace(microsecond=0)
    now_local = now_utc.astimezone(LOCAL_TZ)
    canonical_name = f"BASEIPM_{now_local.strftime('%d%m%Y')}.xlsx"
    raw_path = f"{RAW_PREFIX}{now_local.strftime('%Y/%m')}/{canonical_name}"

    existing = _read_blob(raw_path)
    if existing is not None:
        existing_digest = hashlib.sha256(existing).hexdigest()
        if existing_digest != digest:
            raise IPMUpdateError(
                409,
                f"Já existe uma base para {now_local.strftime('%d/%m/%Y')} com conteúdo diferente. A substituição foi bloqueada.",
            )
        for item in _load_index():
            if item.get("canonical_name") == canonical_name and item.get("metrics", {}).get("unique_protocols") == validation["unique_protocols"]:
                return {"ok": True, "duplicate": True, "run": item}

    if existing is None:
        _write_blob(
            raw_path,
            body,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            overwrite=False,
        )

    run_id = f"{now_local.strftime('%Y%m%dT%H%M%S')}-{digest[:8]}"
    run = {
        "id": run_id,
        "status": "PREPARED",
        "phase": "VALIDATION_COMPLETE",
        "created_at": now_utc.isoformat(),
        "local_created_at": now_local.isoformat(),
        "source_name": _clean(source_name)[:180],
        "canonical_name": canonical_name,
        "raw_path": raw_path,
        "storage": "private-blob",
        "metrics": validation,
        "steps": _steps(),
        "can_publish": False,
        "message": "Base recebida, armazenada e validada. A comparação com o snapshot anterior será iniciada automaticamente.",
    }
    _persist_run(run)
    return {"ok": True, "duplicate": False, "run": run}


__all__ = [
    "IPMUpdateError",
    "create_ipm_import",
    "get_ipm_import",
    "list_ipm_imports",
    "process_ipm_import",
    "validate_ipm_xlsx",
]
