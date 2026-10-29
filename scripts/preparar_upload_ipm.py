"""Prepara um XLSX do IPM para envio manual ao painel, sem modificar a origem.

Uso: python scripts/preparar_upload_ipm.py "C:\\caminho\\relatorio.xlsx"
Requer openpyxl (requirements.txt).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from backend.ipm_update_store import HEADER_ALIASES, REQUIRED_FIELDS, _parse_protocol


def preparar(origem: Path, destino: Path) -> dict:
    origem = origem.expanduser().resolve()
    destino = destino.expanduser().resolve()
    if origem.suffix.lower() != ".xlsx" or not origem.is_file():
        raise ValueError("Informe um arquivo XLSX existente.")
    if origem == destino:
        raise ValueError("Destino não pode substituir a fonte.")
    if origem.stat().st_size > 12 * 1024 * 1024:
        raise ValueError("Arquivo excede 12 MB; verifique o limite de upload do sistema.")
    workbook = load_workbook(origem, read_only=True, data_only=True)
    try:
        sheet = workbook["Report"] if "Report" in workbook.sheetnames else workbook.active
        headers = None
        header_row = 0
        for idx, values in enumerate(sheet.iter_rows(min_row=1, max_row=12, values_only=True), 1):
            found = {str(v).strip(): col for col, v in enumerate(values) if v is not None}
            mapping = {key: next((found[a] for a in aliases if a in found), None)
                       for key, aliases in HEADER_ALIASES.items()}
            if all(mapping.get(k) is not None for k in REQUIRED_FIELDS):
                headers, header_row = mapping, idx
                break
        if headers is None:
            raise ValueError("Cabeçalhos IPM obrigatórios não encontrados nas 12 primeiras linhas.")
        protocols = set()
        invalid = 0
        duplicates = 0
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            col = headers["protocol"]
            value = values[col] if col < len(values) else None
            protocol = _parse_protocol(value)
            if not protocol:
                invalid += 1
            elif protocol in protocols:
                duplicates += 1
            else:
                protocols.add(protocol)
        if invalid or duplicates or not protocols:
            raise ValueError(f"Validação bloqueada: {invalid} protocolos inválidos, {duplicates} duplicados, {len(protocols)} únicos.")
    finally:
        workbook.close()
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(origem, destino)
    digest = hashlib.sha256(destino.read_bytes()).hexdigest()
    return {"arquivo": str(destino), "protocolos_unicos": len(protocols),
            "sha256": digest, "observacao": "Arquivo validado localmente; não foi enviado nem publicado."}


def main() -> None:
    parser = argparse.ArgumentParser(description="Verifica XLSX IPM antes do upload manual.")
    parser.add_argument("origem", type=Path)
    parser.add_argument("--saida", type=Path, default=None)
    args = parser.parse_args()
    destino = args.saida or (Path.home() / "Downloads" / f"BASEIPM_UPLOAD_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    try:
        print(json.dumps(preparar(args.origem, destino), ensure_ascii=False, indent=2))
    except (OSError, ValueError) as exc:
        parser.exit(1, f"ERRO: {exc}\n")


if __name__ == "__main__":
    main()
